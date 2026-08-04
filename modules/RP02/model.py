"""RP02 Finance Reports — Bill Master upload: parsing, customer-master
reconciliation (VCUM01 / vessel_customers), per-financial-year replace.
Mirrors the RP01 historical-data upload pattern."""
import difflib

from database import get_db, get_cursor

TABLE = 'rp02_bill_master'

# Ordered stored fields (one source of truth).
COLUMNS = [
    'financial_year', 'month_label', 'vessel_name', 'material_po',
    'customer_name', 'cargo_type', 'cargo_name', 'bl_qty', 'load_port',
    'mv_mbc', 'discharge_commence', 'discharge_completed',
    'bill_no', 'credit_note', 'old_bill',
]

# CSV header (normalized: BOM stripped, whitespace collapsed, lowercased)
# → stored field. 'Sr' is ignored — display order is regenerated.
HEADER_MAP = {
    'year':                'financial_year',
    'month':               'month_label',
    'm.vessel name':       'vessel_name',
    'vessel name':         'vessel_name',
    'material po':         'material_po',
    'customer name':       'customer_name',
    'type of cargo':       'cargo_type',
    'cargo':               'cargo_name',
    'b/l qty. (mt)':       'bl_qty',
    'b/l qty (mt)':        'bl_qty',
    'load port':           'load_port',
    'mv/mbc':              'mv_mbc',
    'discharge commence':  'discharge_commence',
    'discharge completed': 'discharge_completed',
    'bill no':             'bill_no',
    'credit note':         'credit_note',
    'old bill':            'old_bill',
}

# Pretty headers for the downloadable CSV template (matches the bill master
# file finance already maintains, minus the trailing blank columns).
TEMPLATE_HEADERS = [
    'Sr', 'Year', 'Month', 'M.Vessel Name', 'Material PO', 'Customer Name',
    'Type of Cargo', 'Cargo', 'B/L Qty. (MT)', 'Load  Port', 'MV/MBC',
    'Discharge Commence', 'Discharge Completed', 'Bill No', 'Credit Note', 'Old Bill',
]

TEMPLATE_EXAMPLE_ROWS = [
    ['1', '2026-27', 'Apr-26', 'MV OBE Lotus', '4100550083', 'JSW Steel Ltd',
     'IBRM', 'Orissa Fines', '19,500.00', 'Orissa', 'MV',
     '27-03-2026 20:30', '03-04-2026 05:30', 'DPPL/26-27/12', '', ''],
    ['2', '2026-27', 'Apr-26', 'JSW Devgad', '4300024520', 'JSW Steel Ltd',
     'CBRM', 'PSH Coal', '7,916.00', 'Jaigad', 'MBC',
     '02-04-2026 01:45', '02-04-2026 10:05', 'DPPL/26-27/13', '', ''],
]


# ── Pure helpers ─────────────────────────────────────────────────────────────
def _blank(v):
    return v is None or (isinstance(v, str) and v.strip() == '')


def _norm_header(h):
    return ' '.join(str(h or '').replace('﻿', '').split()).lower()


def parse_number(v):
    """Numeric → float ('" 37,930.00 "' → 37930.0). Blank→None. Else ValueError."""
    if _blank(v):
        return None
    try:
        return float(str(v).strip().replace(',', ''))
    except (TypeError, ValueError):
        raise ValueError(f"invalid number '{v}'")


def suggest_matches(value, master_values, n=3, cutoff=0.6):
    """Closest master values to `value` (case-insensitive), up to n."""
    if not value or not master_values:
        return []
    lower_map = {}
    for m in master_values:
        lower_map.setdefault(str(m).lower(), m)
    hits = difflib.get_close_matches(str(value).lower(), list(lower_map.keys()), n=n, cutoff=cutoff)
    return [lower_map[h] for h in hits]


def parse_rows(headers, raw_rows):
    """Map raw spreadsheet rows to field dicts using normalized `headers`.

    Returns (rows, errors). A row is skipped if every cell is blank.
    Required: financial_year, vessel_name, customer_name. bl_qty must be
    numeric; mv_mbc must be MV or MBC when present. Discharge columns are
    kept as raw text (prod data mixes timestamps with 'Under Discharge').
    Row numbers are 1-based spreadsheet rows (header row included)."""
    idx = {}
    for i, h in enumerate(headers):
        field = HEADER_MAP.get(_norm_header(h))
        if field and field not in idx:
            idx[field] = i

    rows, errors = [], []

    def cell(r, key):
        i = idx.get(key)
        return r[i] if (i is not None and i < len(r)) else None

    for n, r in enumerate(raw_rows, start=2):  # +1 header, +1 to 1-base
        if all(_blank(c) for c in r):
            continue
        rec, row_errs = {}, []
        for key in COLUMNS:
            if key == 'bl_qty':
                continue
            v = cell(r, key)
            rec[key] = (str(v).strip() if not _blank(v) else None)
        try:
            rec['bl_qty'] = parse_number(cell(r, 'bl_qty'))
        except ValueError as e:
            row_errs.append(str(e))

        if _blank(rec.get('financial_year')):
            row_errs.append('Year is required')
        if _blank(rec.get('vessel_name')):
            row_errs.append('M.Vessel Name is required')
        if _blank(rec.get('customer_name')):
            row_errs.append('Customer Name is required')
        if rec.get('mv_mbc'):
            mv = rec['mv_mbc'].upper()
            if mv not in ('MV', 'MBC'):
                row_errs.append(f"MV/MBC must be 'MV' or 'MBC' (got '{rec['mv_mbc']}')")
            else:
                rec['mv_mbc'] = mv

        if row_errs:
            errors.append({'row': n, 'message': '; '.join(row_errs)})
        else:
            rows.append(rec)
    return rows, errors


def parse_upload(file_storage):
    """Parse a werkzeug FileStorage (.csv/.xlsx) → (rows, errors).
    The header row is the first row containing both 'Customer Name' and
    'Bill No' (normalized), so leading title rows are tolerated.
    The format is detected from the CONTENT, not the extension — a CSV
    saved with any extension still parses, and unreadable files come back
    as a friendly format error instead of a 500."""
    import io, csv as _csv
    raw = file_storage.read()

    def from_matrix(matrix):
        header_idx = None
        for i, row in enumerate(matrix):
            normed = {_norm_header(c) for c in row}
            if 'customer name' in normed and 'bill no' in normed:
                header_idx = i
                break
        if header_idx is None:
            return [], [{'row': 0, 'message': "Could not find a header row containing 'Customer Name' and 'Bill No'"}]
        headers = [str(c).strip() if c is not None else '' for c in matrix[header_idx]]
        return parse_rows(headers, matrix[header_idx + 1:])

    if raw[:2] == b'PK':  # zip container → real .xlsx
        try:
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
            ws = wb.active
            matrix = [list(r) for r in ws.iter_rows(values_only=True)]
        except Exception:
            return [], [{'row': 0, 'message': 'Could not read the Excel file — re-save it as .xlsx or CSV and try again'}]
        return from_matrix(matrix)
    if raw[:4] == b'\xd0\xcf\x11\xe0':  # OLE2 container → legacy .xls
        return [], [{'row': 0, 'message': "Old Excel '.xls' format is not supported — save the file as .xlsx or CSV"}]
    # Anything else is treated as CSV text regardless of extension.
    text = raw.decode('utf-8-sig', errors='replace')
    matrix = list(_csv.reader(io.StringIO(text)))
    return from_matrix(matrix)


def build_template_csv():
    """Return the upload template as CSV text (headers + example rows)."""
    import io, csv as _csv
    buf = io.StringIO()
    w = _csv.writer(buf)
    w.writerow(TEMPLATE_HEADERS)
    for r in TEMPLATE_EXAMPLE_ROWS:
        w.writerow(r)
    return buf.getvalue()


# ── Customer master reconciliation (VCUM01 / vessel_customers) ──────────────
def get_customer_master():
    conn = get_db()
    cur = get_cursor(conn)
    try:
        cur.execute('SELECT name FROM vessel_customers ORDER BY name')
        return [r['name'] for r in cur.fetchall() if r['name'] not in (None, '')]
    finally:
        conn.close()


def reconcile(rows, customers):
    """Split the rows' distinct customer names into recognized vs unknown
    (with fuzzy suggestions from the customer master).
    Returns {'customer_name': {recognized:[...], unknown:[{value,count,suggestions}]}}."""
    counts = {}
    for r in rows:
        v = r.get('customer_name')
        if v:
            counts[v] = counts.get(v, 0) + 1
    valid = {str(c).lower() for c in customers}
    recognized, unknown = [], []
    for value, count in sorted(counts.items()):
        if str(value).lower() in valid:
            recognized.append(value)
        else:
            unknown.append({'value': value, 'count': count,
                            'suggestions': suggest_matches(value, customers)})
    return {'customer_name': {'recognized': recognized, 'unknown': unknown}}


def apply_resolutions(rows, resolutions):
    """Apply 'replace' resolutions to parsed rows (pure). `resolutions` is
    {column: {old_value: {'action': 'replace'|'add'|'keep', 'target': new}}}.
    Only 'replace' with a non-empty target rewrites the cell."""
    repl = {}
    for col, mapping in (resolutions or {}).items():
        for old, res in (mapping or {}).items():
            if isinstance(res, dict) and res.get('action') == 'replace' and res.get('target'):
                repl.setdefault(col, {})[old] = res['target']
    if not repl:
        return [dict(r) for r in rows]
    out = []
    for r in rows:
        nr = dict(r)
        for col, m in repl.items():
            if nr.get(col) in m:
                nr[col] = m[nr[col]]
        out.append(nr)
    return out


def add_customer(name):
    """Insert `name` into vessel_customers if not already there (case-insensitive).
    Returns True if inserted. Master details (GSTIN, SAP code…) are filled in
    later via VCUM01."""
    conn = get_db()
    cur = get_cursor(conn)
    try:
        cur.execute(
            "INSERT INTO vessel_customers (name) "
            "SELECT %s WHERE NOT EXISTS "
            "(SELECT 1 FROM vessel_customers WHERE LOWER(name) = LOWER(%s))",
            [name, name],
        )
        inserted = cur.rowcount > 0
        conn.commit()
        return inserted
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


# ── Storage ──────────────────────────────────────────────────────────────────
def replace_years(rows, uploaded_by):
    """Replace all stored rows for the financial year(s) present in the file,
    in one transaction. Other years are untouched. Returns (inserted, [years])."""
    years = sorted({r['financial_year'] for r in rows})
    conn = get_db()
    cur = get_cursor(conn)
    try:
        if years:
            placeholders = ', '.join(['%s'] * len(years))
            cur.execute(f"DELETE FROM {TABLE} WHERE financial_year IN ({placeholders})", years)
        insert_cols = COLUMNS + ['uploaded_by']
        sql = (f"INSERT INTO {TABLE} ({', '.join(insert_cols)}) "
               f"VALUES ({', '.join(['%s'] * len(insert_cols))})")
        for r in rows:
            cur.execute(sql, [r.get(c) for c in COLUMNS] + [uploaded_by])
        conn.commit()
        return len(rows), years
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


def get_status():
    """Return {count, uploaded_at, years} for the current dataset."""
    conn = get_db()
    cur = get_cursor(conn)
    try:
        cur.execute(f"SELECT COUNT(*) AS c, MAX(uploaded_at) AS at FROM {TABLE}")
        r = cur.fetchone()
        cur.execute(f"SELECT DISTINCT financial_year FROM {TABLE} ORDER BY financial_year")
        years = [x['financial_year'] for x in cur.fetchall()]
        at = r['at']
        return {'count': r['c'], 'uploaded_at': at.isoformat() if at else None,
                'years': years}
    finally:
        conn.close()


_CAST_FILTER_COLS = {'bl_qty'}


def get_rows(page=1, size=50, filters=None):
    """Paginated stored rows. `filters` is a list of {field, value}; each is
    ILIKE-matched against the FULL dataset and AND-combined."""
    conn = get_db()
    cur = get_cursor(conn)
    try:
        allowed = set(COLUMNS)
        where_parts, params = [], []
        for f in (filters or []):
            field = f.get('field')
            val = (f.get('value') or '').strip() if f.get('value') is not None else ''
            if field in allowed and val:
                col = f"{field}::TEXT" if field in _CAST_FILTER_COLS else field
                where_parts.append(f"{col} ILIKE %s")
                params.append(f"%{val}%")
        where = ('WHERE ' + ' AND '.join(where_parts)) if where_parts else ''
        cur.execute(f"SELECT COUNT(*) AS c FROM {TABLE} {where}", params)
        total = cur.fetchone()['c']
        offset = (page - 1) * size
        cur.execute(
            f"SELECT * FROM {TABLE} {where} ORDER BY financial_year, id LIMIT %s OFFSET %s",
            params + [size, offset])
        rows = []
        for r in cur.fetchall():
            d = dict(r)
            if d.get('bl_qty') is not None:
                d['bl_qty'] = float(d['bl_qty'])
            if d.get('uploaded_at') is not None:
                d['uploaded_at'] = str(d['uploaded_at'])
            rows.append(d)
        return rows, total
    finally:
        conn.close()


def update_row(row_id, data):
    """Validate + update a single stored row. Returns {'success': True} or
    {'error': msg}. Mirrors the upload validators."""
    clean = {}
    try:
        clean['bl_qty'] = parse_number(data.get('bl_qty'))
    except ValueError as e:
        return {'error': str(e)}
    for k in COLUMNS:
        if k == 'bl_qty':
            continue
        v = data.get(k)
        clean[k] = (str(v).strip() if v not in (None, '') else None)
    for req, label in (('financial_year', 'Year'), ('vessel_name', 'M.Vessel Name'),
                       ('customer_name', 'Customer Name')):
        if not clean.get(req):
            return {'error': f'{label} is required'}
    if clean.get('mv_mbc'):
        mv = clean['mv_mbc'].upper()
        if mv not in ('MV', 'MBC'):
            return {'error': "MV/MBC must be 'MV' or 'MBC'"}
        clean['mv_mbc'] = mv

    conn = get_db()
    cur = get_cursor(conn)
    try:
        sets = ', '.join([f"{c}=%s" for c in COLUMNS])
        cur.execute(f"UPDATE {TABLE} SET {sets} WHERE id=%s",
                    [clean.get(c) for c in COLUMNS] + [row_id])
        conn.commit()
        return {'success': True}
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


def delete_row(row_id):
    conn = get_db()
    cur = get_cursor(conn)
    try:
        cur.execute(f"DELETE FROM {TABLE} WHERE id=%s", [row_id])
        conn.commit()
    finally:
        conn.close()


# ── Bill Master Report (uploaded rows + live invoice-derived rows) ──────────

# ponytail: hardcoded billing go-live boundary — the uploaded bill master
# covers everything before invoice DPPL/26-27/158; rows from that invoice
# onward are derived live from invoice/cargo data. Bump if the boundary moves.
LIVE_FROM_FY = '2026-27'
LIVE_FROM_SEQ = 158

# One row per invoiced cargo declaration line, shaped like the bill master.
# Discharge timings reuse the RP01 cargo-report logic (ldud_anchorage /
# mbc_discharge_port_lines); Material PO comes from the latest LDUD per VCN
# (vessels) or the MBC customer detail line (MBCs).
_LIVE_ROWS_SQL = """
WITH inv AS (
    SELECT id, invoice_number, financial_year, invoice_date, customer_name
    FROM invoice_header
    WHERE invoice_status <> 'Cancelled'
      AND (financial_year > %(fy)s
           OR (financial_year = %(fy)s AND COALESCE(doc_series_seq, 0) >= %(seq)s))
),
cn AS (
    SELECT original_invoice_number, STRING_AGG(doc_number, ', ') AS credit_notes
    FROM fdcn_header
    WHERE doc_type = 'CN' AND COALESCE(doc_status, '') <> 'Cancelled'
      AND COALESCE(original_invoice_number, '') <> ''
    GROUP BY original_invoice_number
),
ldud_latest AS (
    SELECT DISTINCT ON (vcn_id) vcn_id, id AS ldud_id, material_po_number
    FROM ldud_header
    ORDER BY vcn_id, id DESC
),
ldud_times AS (
    SELECT ldud_id,
           MIN(discharge_started) AS commence,
           CASE WHEN SUM(CASE WHEN discharge_started IS NOT NULL
                               AND discharge_commenced IS NULL THEN 1 ELSE 0 END) > 0
                THEN NULL ELSE MAX(discharge_commenced) END AS completed
    FROM ldud_anchorage
    GROUP BY ldud_id
),
mbc_times AS (
    SELECT mbc_id,
           MIN(NULLIF(TRIM(unloading_commenced), '')::timestamp) AS commence,
           MAX(NULLIF(TRIM(unloading_completed), '')::timestamp) AS completed
    FROM mbc_discharge_port_lines
    GROUP BY mbc_id
)
SELECT inv.financial_year, inv.invoice_date, inv.invoice_number, inv.customer_name,
       vh.vessel_name, ll.material_po_number AS material_po,
       vh.cargo_type, cd.cargo_name, cd.bl_quantity AS bl_qty, vh.load_port,
       'MV' AS mv_mbc,
       lt.commence::text AS discharge_commence,
       lt.completed::text AS discharge_completed,
       cn.credit_notes
FROM inv
JOIN invoice_bill_mapping ibm ON ibm.invoice_id = inv.id
JOIN bill_lines bl ON bl.bill_id = ibm.bill_id AND bl.cargo_source_type = 'VCN_IMPORT'
JOIN vcn_cargo_declaration cd ON cd.id = bl.cargo_source_id
JOIN vcn_header vh ON vh.id = cd.vcn_id
LEFT JOIN ldud_latest ll ON ll.vcn_id = cd.vcn_id
LEFT JOIN ldud_times lt ON lt.ldud_id = ll.ldud_id
LEFT JOIN cn ON cn.original_invoice_number = inv.invoice_number

UNION ALL

SELECT inv.financial_year, inv.invoice_date, inv.invoice_number, inv.customer_name,
       vh.vessel_name, ll.material_po_number AS material_po,
       vh.cargo_type, cd.cargo_name, cd.bl_quantity AS bl_qty, vh.load_port,
       'MV' AS mv_mbc,
       lt.commence::text AS discharge_commence,
       lt.completed::text AS discharge_completed,
       cn.credit_notes
FROM inv
JOIN invoice_bill_mapping ibm ON ibm.invoice_id = inv.id
JOIN bill_lines bl ON bl.bill_id = ibm.bill_id AND bl.cargo_source_type = 'VCN_EXPORT'
JOIN vcn_export_cargo_declaration cd ON cd.id = bl.cargo_source_id
JOIN vcn_header vh ON vh.id = cd.vcn_id
LEFT JOIN ldud_latest ll ON ll.vcn_id = cd.vcn_id
LEFT JOIN ldud_times lt ON lt.ldud_id = ll.ldud_id
LEFT JOIN cn ON cn.original_invoice_number = inv.invoice_number

UNION ALL

SELECT inv.financial_year, inv.invoice_date, inv.invoice_number, inv.customer_name,
       mh.mbc_name AS vessel_name, cd.material_po,
       mh.cargo_type, cd.cargo_name, cd.quantity AS bl_qty, mh.load_port,
       'MBC' AS mv_mbc,
       mt.commence::text AS discharge_commence,
       mt.completed::text AS discharge_completed,
       cn.credit_notes
FROM inv
JOIN invoice_bill_mapping ibm ON ibm.invoice_id = inv.id
JOIN bill_lines bl ON bl.bill_id = ibm.bill_id AND bl.cargo_source_type = 'MBC'
JOIN mbc_customer_details cd ON cd.id = bl.cargo_source_id
JOIN mbc_header mh ON mh.id = cd.mbc_id
LEFT JOIN mbc_times mt ON mt.mbc_id = cd.mbc_id
LEFT JOIN cn ON cn.original_invoice_number = inv.invoice_number

ORDER BY financial_year, invoice_number, vessel_name
"""


def _month_label(d):
    """invoice_date → 'Apr-26' style label (matches the bill master CSV)."""
    if not d:
        return None
    from datetime import datetime as _dt
    try:
        return _dt.strptime(str(d)[:10], '%Y-%m-%d').strftime('%b-%y')
    except ValueError:
        return None


def _fmt_ts(ts):
    """DB timestamp text → 'DD-MM-YYYY HH:MM' (matches the bill master CSV)."""
    if not ts:
        return None
    from datetime import datetime as _dt
    try:
        return _dt.strptime(str(ts).replace('T', ' ')[:16], '%Y-%m-%d %H:%M').strftime('%d-%m-%Y %H:%M')
    except ValueError:
        return str(ts)


def get_bill_master_report():
    """The bill-master-shaped report: all uploaded rows as-is (file order),
    followed by rows derived from actual invoice data from
    DPPL/{LIVE_FROM_FY}/{LIVE_FROM_SEQ} onward. Each row carries a `source`
    of 'Uploaded' or 'Live'."""
    conn = get_db()
    cur = get_cursor(conn)
    try:
        cur.execute(f"SELECT * FROM {TABLE} ORDER BY id")
        out = []
        for r in cur.fetchall():
            d = {c: r.get(c) for c in COLUMNS}
            if d.get('bl_qty') is not None:
                d['bl_qty'] = float(d['bl_qty'])
            d['source'] = 'Uploaded'
            out.append(d)

        cur.execute(_LIVE_ROWS_SQL, {'fy': LIVE_FROM_FY, 'seq': LIVE_FROM_SEQ})
        for r in cur.fetchall():
            out.append({
                'financial_year':      r['financial_year'],
                'month_label':         _month_label(r['invoice_date']),
                'vessel_name':         r['vessel_name'],
                'material_po':         r['material_po'],
                'customer_name':       r['customer_name'],
                'cargo_type':          r['cargo_type'],
                'cargo_name':          r['cargo_name'],
                'bl_qty':              float(r['bl_qty']) if r['bl_qty'] is not None else None,
                'load_port':           r['load_port'],
                'mv_mbc':              r['mv_mbc'],
                'discharge_commence':  _fmt_ts(r['discharge_commence']),
                'discharge_completed': _fmt_ts(r['discharge_completed']),
                'bill_no':             r['invoice_number'],
                'credit_note':         r['credit_notes'],
                'old_bill':            None,
                'source':              'Live',
            })
        return out
    finally:
        conn.close()


# ── Billing Pipeline dashboard ───────────────────────────────────────────────
# The bill master answers "what has been billed". A billing desk also needs the
# complement: cargo discharged but not yet on an invoice. Same row shape, plus a
# status, so both halves sit in one grid and one export.
#
# Deliberately quantity-only (MT) and count-only — no rates, no amounts. The
# question this answers is "what is waiting on me", not "what is it worth".

# Cargo with no invoice behind it yet, in bill-master shape. `ready` mirrors the
# FIN01 gate: parent document open enough to bill. Not-ready rows are blocked
# upstream in operations, not by finance.
_PENDING_ROWS_SQL = """
WITH ldud_latest AS (
    SELECT DISTINCT ON (vcn_id) vcn_id, id AS ldud_id, material_po_number, doc_status
    FROM ldud_header
    ORDER BY vcn_id, id DESC
),
ldud_times AS (
    SELECT ldud_id,
           MIN(discharge_started) AS commence,
           CASE WHEN SUM(CASE WHEN discharge_started IS NOT NULL
                               AND discharge_commenced IS NULL THEN 1 ELSE 0 END) > 0
                THEN NULL ELSE MAX(discharge_commenced) END AS completed
    FROM ldud_anchorage
    GROUP BY ldud_id
),
mbc_times AS (
    SELECT mbc_id,
           MIN(NULLIF(TRIM(unloading_commenced), '')::timestamp) AS commence,
           MAX(NULLIF(TRIM(unloading_completed), '')::timestamp) AS completed
    FROM mbc_discharge_port_lines
    GROUP BY mbc_id
)
SELECT vh.vessel_name, ll.material_po_number AS material_po, cd.customer_name,
       vh.cargo_type, cd.cargo_name,
       COALESCE(cd.bl_quantity, 0) - COALESCE(cd.billed_quantity, 0) AS pending_qty,
       vh.load_port, 'MV' AS mv_mbc,
       lt.commence::text  AS discharge_commence,
       lt.completed::text AS discharge_completed,
       COALESCE(ll.doc_status, 'No LDUD') AS doc_status,
       CASE WHEN ll.doc_status IN ('Closed', 'Partial Close') THEN 1 ELSE 0 END AS ready
FROM vcn_cargo_declaration cd
JOIN vcn_header vh ON vh.id = cd.vcn_id
LEFT JOIN ldud_latest ll ON ll.vcn_id = cd.vcn_id
LEFT JOIN ldud_times  lt ON lt.ldud_id = ll.ldud_id
WHERE (COALESCE(cd.is_billed, 0) = 0 OR COALESCE(cd.billed_quantity, 0) < cd.bl_quantity)

UNION ALL

SELECT vh.vessel_name, ll.material_po_number, cd.customer_name,
       vh.cargo_type, cd.cargo_name,
       COALESCE(cd.bl_quantity, 0) - COALESCE(cd.billed_quantity, 0),
       vh.load_port, 'MV',
       lt.commence::text, lt.completed::text,
       COALESCE(ll.doc_status, 'No LDUD'),
       CASE WHEN ll.doc_status IN ('Closed', 'Partial Close') THEN 1 ELSE 0 END
FROM vcn_export_cargo_declaration cd
JOIN vcn_header vh ON vh.id = cd.vcn_id
LEFT JOIN ldud_latest ll ON ll.vcn_id = cd.vcn_id
LEFT JOIN ldud_times  lt ON lt.ldud_id = ll.ldud_id
WHERE (COALESCE(cd.is_billed, 0) = 0 OR COALESCE(cd.billed_quantity, 0) < cd.bl_quantity)

UNION ALL

SELECT mh.mbc_name, cd.material_po, cd.customer_name,
       mh.cargo_type, cd.cargo_name,
       COALESCE(cd.quantity, 0) - COALESCE(cd.billed_quantity, 0),
       mh.load_port, 'MBC',
       mt.commence::text, mt.completed::text,
       COALESCE(mh.doc_status, ''),
       CASE WHEN mh.doc_status IN ('Approved', 'Closed', 'Partial Close') THEN 1 ELSE 0 END
FROM mbc_customer_details cd
JOIN mbc_header mh ON mh.id = cd.mbc_id
LEFT JOIN mbc_times mt ON mt.mbc_id = cd.mbc_id
WHERE (COALESCE(cd.is_billed, 0) = 0 OR COALESCE(cd.billed_quantity, 0) < cd.quantity)
"""

# Days-since-discharge buckets, oldest last so the ordinal colour ramp and the
# bar order agree.
AGE_BUCKETS = [('0-7 days', 0, 7), ('8-15 days', 8, 15),
               ('16-30 days', 16, 30), ('30+ days', 31, None)]

BUCKET_LABELS = [b[0] for b in AGE_BUCKETS] + ['Undated']


def _age_days(completed):
    """Whole days between a discharge-completed stamp and now. None if unparseable."""
    from datetime import datetime as _dt
    if not completed:
        return None
    txt = str(completed).replace('T', ' ').strip()
    for fmt, width in (('%d-%m-%Y %H:%M', 16), ('%Y-%m-%d %H:%M', 16), ('%Y-%m-%d', 10)):
        try:
            return max((_dt.now() - _dt.strptime(txt[:width], fmt)).days, 0)
        except ValueError:
            continue
    return None


def _bucket(days):
    if days is None:
        return 'Undated'
    for label, lo, hi in AGE_BUCKETS:
        if days >= lo and (hi is None or days <= hi):
            return label
    return 'Undated'


def get_pending_rows():
    """Cargo declared/discharged but not yet invoiced, in bill-master shape."""
    conn = get_db()
    cur = get_cursor(conn)
    try:
        cur.execute(_PENDING_ROWS_SQL)
        out = []
        for r in cur.fetchall():
            qty = float(r['pending_qty'] or 0)
            if qty <= 0:
                continue                     # nothing left to bill on this line
            days = _age_days(_fmt_ts(r['discharge_completed']))
            out.append({
                'vessel_name':         r['vessel_name'],
                'material_po':         r['material_po'],
                'customer_name':       r['customer_name'],
                'cargo_type':          r['cargo_type'],
                'cargo_name':          r['cargo_name'],
                'bl_qty':              round(qty, 2),
                'load_port':           r['load_port'],
                'mv_mbc':              r['mv_mbc'],
                'discharge_commence':  _fmt_ts(r['discharge_commence']),
                'discharge_completed': _fmt_ts(r['discharge_completed']),
                'doc_status':          r['doc_status'],
                'status':              'Ready to bill' if r['ready'] else 'Blocked upstream',
                'age_days':            days,
                'age_bucket':          _bucket(days),
            })
        return out
    finally:
        conn.close()


def _top(rows, key, n=8):
    """Top-n {label, lines, qty} by pending quantity; the tail folds into Other."""
    agg = {}
    for r in rows:
        k = (r.get(key) or '').strip() or '—'
        a = agg.setdefault(k, {'label': k, 'lines': 0, 'qty': 0.0})
        a['lines'] += 1
        a['qty'] += r['bl_qty']
    ordered = sorted(agg.values(), key=lambda a: -a['qty'])
    if len(ordered) > n:
        head, tail = ordered[:n], ordered[n:]
        ordered = head + [{'label': f'Other ({len(tail)})',
                           'lines': sum(a['lines'] for a in tail),
                           'qty': sum(a['qty'] for a in tail)}]
    return [{**a, 'qty': round(a['qty'], 2)} for a in ordered]


def get_billing_dashboard():
    """Everything the billing pipeline dashboard renders.

    Quantities (MT) and line counts only — no rates and no amounts anywhere in
    this payload, by design."""
    from datetime import datetime as _dt

    pending = get_pending_rows()
    ready = [r for r in pending if r['status'] == 'Ready to bill']
    blocked = [r for r in pending if r['status'] == 'Blocked upstream']

    billed = get_bill_master_report()
    this_month = _dt.now().strftime('%b-%y')

    def qty(rows):
        return round(sum(r['bl_qty'] or 0 for r in rows), 2)

    buckets = {label: {'label': label, 'lines': 0, 'qty': 0.0} for label in BUCKET_LABELS}
    for r in ready:
        b = buckets[r['age_bucket']]
        b['lines'] += 1
        b['qty'] += r['bl_qty']

    aged = [r['age_days'] for r in ready if r['age_days'] is not None]

    return {
        'generated_at': _dt.now().strftime('%d-%m-%Y %H:%M'),
        'kpi': {
            'ready_lines':   len(ready),
            'ready_qty':     qty(ready),
            'blocked_lines': len(blocked),
            'blocked_qty':   qty(blocked),
            'oldest_days':   max(aged) if aged else 0,
            'customers':     len({(r['customer_name'] or '—') for r in ready}),
            'billed_lines':  len(billed),
            'billed_month':  len([r for r in billed if r.get('month_label') == this_month]),
            'month_label':   this_month,
        },
        'ageing':            [{**buckets[l], 'qty': round(buckets[l]['qty'], 2)}
                              for l in BUCKET_LABELS],
        'by_customer':       _top(ready, 'customer_name'),
        'by_cargo':          _top(ready, 'cargo_type', n=6),
        'by_mode':           _top(pending, 'mv_mbc', n=4),
        'blocked_by_status': _top(blocked, 'doc_status', n=6),
        'rows':              sorted(pending,
                                    key=lambda r: (-(r['age_days'] if r['age_days'] is not None else -1),
                                                   r['customer_name'] or '')),
    }

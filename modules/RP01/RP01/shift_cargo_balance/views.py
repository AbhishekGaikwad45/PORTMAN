from flask import render_template, request, jsonify, session, redirect, url_for, Response
from functools import wraps
from datetime import date, datetime, timedelta
import io
import json
import re

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from .. import bp
from database import get_db, get_cursor

REPORT_CUTOFF_DATE = date(2026, 9, 1)
REPORT_CUTOFF_DT = datetime(2026, 9, 1, 0, 0, 0)


def login_required(f):
    @wraps(f)
    def decorated(*args, **kwargs):
        if 'user_id' not in session:
            return redirect(url_for('login'))
        return f(*args, **kwargs)
    return decorated


def safe_dt(value):
    if not value:
        return None
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime.combine(value, datetime.min.time())

    value = str(value).strip()
    if not value or value.lower() in ('none', 'null', '—', '-', ''):
        return None

    try:
        if 'T' in value or len(value) == 10:
            return datetime.fromisoformat(value)
    except Exception:
        pass

    formats = [
        "%Y-%m-%d %H:%M:%S",
        "%Y-%m-%d %H:%M",
        "%d-%m-%Y %H:%M:%S",
        "%d-%m-%Y %H:%M",
        "%Y-%m-%dT%H:%M:%S",
        "%Y-%m-%dT%H:%M"
    ]
    for fmt in formats:
        try:
            return datetime.strptime(value, fmt)
        except Exception:
            pass

    return None


def _fmt_dt(val):
    if not val:
        return ''
    d = safe_dt(val)
    return d.strftime('%d-%m-%Y %H:%M') if d else ''


def _clean_cargo_name(raw):
    if not raw:
        return 'General Cargo'
    c = str(raw).strip()
    return re.sub(r'\s+', ' ', c)


def _is_wr19(berth_str):
    if not berth_str:
        return False
    u = str(berth_str).strip().upper()
    if '-' in u and ('T' in u or ':' in u):
        return False
    return ('WR' in u and '19' in u) or u in ('WR 19', 'WR19', 'R19', 'R-19', 'WR-19', 'WT @ R19', 'WT R19')


def _is_berth_10_to_12(berth_str):
    if not berth_str:
        return False, ''
    s = str(berth_str).strip().upper()
    if '-' in s and ('T' in s or ':' in s):
        return False, ''
    m = re.search(r'\b(?:BERTH\s*(?:NO\.?)?\s*|B\s*-?\s*)(10|11|12)\b', s)
    if m:
        return True, m.group(1)
    if s in ('10', '11', '12'):
        return True, s
    return False, ''


SHIFT_RANK = {'A': 1, 'B': 2, 'C': 3}


def _is_upcoming_shift(target_date, shift_key, now=None):
    now = now or datetime.now()
    cur_op_date = _operational_date(now)
    cur_shift = _current_shift_code(now)
    if target_date > cur_op_date:
        return True
    if target_date == cur_op_date:
        return SHIFT_RANK.get(shift_key, 0) > SHIFT_RANK.get(cur_shift, 0)
    return False


def _fetch_shift_cargo_balance(report_date_str, shift_key):
    """
    Dynamically computes the cargo balance at the jetty for the given date and shift:
    - Cutoff Date: REPORT_CUTOFF_DATE (01-09-2026). Dates before this are blocked.
    - Upcoming shift data is blocked and not shown.
    - Uses exact/carried BPR CTE query with cutoff date applied directly in SQL.
    - Standard Jetty cargos (Berths 1-9 / Jetty) are aggregated by cargo name.
    - Berth 10, 11, and 12 are separated as (Berth No.X).
    - WR 19 is separated as (WR 19).
    """
    shift_key = (shift_key or 'C').strip().upper()
    if shift_key not in ('A', 'B', 'C'):
        shift_key = 'C'

    try:
        target_date = datetime.strptime(report_date_str, '%Y-%m-%d').date()
    except Exception:
        target_date = datetime.now().date()
    target_date_str = target_date.strftime('%Y-%m-%d')
    cutoff_date_str = REPORT_CUTOFF_DATE.strftime('%Y-%m-%d')
    cutoff_fmt = REPORT_CUTOFF_DATE.strftime('%d-%m-%Y')

    # Cutoff date validation (01-09-2026)
    if target_date < REPORT_CUTOFF_DATE:
        return {
            'entry_date': target_date_str,
            'shift': shift_key,
            'shift_display': f"{shift_key} Shift",
            'title': f"Cargo Balance at Jetty for {shift_key} Shift",
            'items': [],
            'total_qty': 0,
            'total_balance': 0,
            'is_cutoff': True,
            'sms_text': (
                f"Cargo Balance at Jetty for {shift_key} Shift\n\n"
                f"Cutoff Date: {cutoff_fmt}.\n"
                "Reports before this date are not available.\n\n"
                "Total: 0 MT.\n\n"
                "Regards"
            )
        }

    now = datetime.now()
    cur_op_date = _operational_date(now)
    cur_shift = _current_shift_code(now)

    is_upcoming = _is_upcoming_shift(target_date, shift_key, now)

    # If shift is in the future, do not show any upcoming data
    if is_upcoming:
        return {
            'entry_date': target_date_str,
            'shift': shift_key,
            'shift_display': f"{shift_key} Shift",
            'title': f"Cargo Balance at Jetty for {shift_key} Shift",
            'items': [],
            'total_qty': 0,
            'total_balance': 0,
            'sms_text': (
                f"Cargo Balance at Jetty for {shift_key} Shift\n\n"
                "No balance data for upcoming shift.\n\n"
                "Total: 0 MT.\n\n"
                "Regards"
            )
        }

    conn = get_db()
    cur = get_cursor(conn)

    sql_query = r"""
        WITH params AS (
            SELECT
                %s::date AS report_date,
                %s::text AS report_shift,
                %s::date AS cutoff_date
        ),

        bpr AS (
            SELECT
                bpr.berth_layout,
                bpr.shift,
                bpr.report_date,
                bpr.updated_at
            FROM barge_position_report bpr
            CROSS JOIN params p
            WHERE bpr.report_date = p.report_date
              AND bpr.report_date >= p.cutoff_date
              AND bpr.shift = p.report_shift
            ORDER BY bpr.updated_at DESC
            LIMIT 1
        ),

        carried_bpr AS (
            SELECT
                bpr.berth_layout,
                bpr.shift,
                bpr.report_date,
                bpr.updated_at
            FROM barge_position_report bpr
            CROSS JOIN params p
            WHERE NOT EXISTS (
                SELECT 1 FROM bpr
            )
              AND bpr.report_date >= p.cutoff_date
              AND (
                    bpr.report_date < p.report_date
                    OR (
                        bpr.report_date = p.report_date
                        AND CASE bpr.shift
                            WHEN 'C' THEN 3
                            WHEN 'B' THEN 2
                            WHEN 'A' THEN 1
                            ELSE 0
                        END
                        <=
                        CASE p.report_shift
                            WHEN 'C' THEN 3
                            WHEN 'B' THEN 2
                            WHEN 'A' THEN 1
                            ELSE 0
                        END
                    )
                  )
            ORDER BY
                bpr.report_date DESC,
                CASE bpr.shift
                    WHEN 'C' THEN 3
                    WHEN 'B' THEN 2
                    WHEN 'A' THEN 1
                    ELSE 0
                END DESC,
                bpr.updated_at DESC
            LIMIT 1
        ),

        selected_bpr AS (
            SELECT * FROM bpr

            UNION ALL

            SELECT * FROM carried_bpr
        ),

        items AS (
            SELECT
                x.item
            FROM selected_bpr s
            CROSS JOIN LATERAL jsonb_array_elements(
                CASE
                    WHEN jsonb_typeof(s.berth_layout::jsonb) = 'array'
                    THEN s.berth_layout::jsonb
                    ELSE '[]'::jsonb
                END
            ) x(item)
        ),

        candidate AS (
            SELECT
                item->>'id' AS id,
                TRIM(item->>'name') AS name,
                TRIM(item->>'cargo') AS cargo,
                UPPER(TRIM(item->>'berth')) AS berth,
                COALESCE(
                    NULLIF(item->>'balance', '')::numeric,
                    NULLIF(item->>'balance_qty', '')::numeric,
                    0
                ) AS balance
            FROM items
            WHERE COALESCE(TRIM(item->>'name'), '') <> ''
              AND UPPER(TRIM(item->>'berth')) NOT IN
                  ('WAITING', 'NONE', 'NULL', '—', '-', 'EMPTY', '')
        ),

        positive AS (
            SELECT *
            FROM candidate
            WHERE balance > 0.01
        ),

        final_report AS (

            SELECT
                'JETTY' AS berth,
                cargo,
                ROUND(SUM(balance))::numeric AS balance
            FROM positive
            WHERE NOT (
                berth ~ '\m(BERTH[ ]*NO[.]?[ ]*|B[ ]*-?)(10|11|12)\M'
                OR berth IN ('10','11','12')
            )
            AND NOT (
                berth ~ '\m(WR[ ]*19|WR-19|R-?19)\M'
                OR berth LIKE '%%WR%%19%%'
            )
            GROUP BY cargo

            UNION ALL

            SELECT
                berth,
                cargo,
                ROUND(SUM(balance))::numeric AS balance
            FROM positive
            WHERE (
                berth ~ '\m(BERTH[ ]*NO[.]?[ ]*|B[ ]*-?)(10|11|12)\M'
                OR berth IN ('10','11','12')
            )
            GROUP BY berth, cargo

            UNION ALL

            SELECT
                'WR 19' AS berth,
                cargo,
                ROUND(SUM(balance))::numeric AS balance
            FROM positive
            WHERE (
                berth ~ '\m(WR[ ]*19|WR-19|R-?19)\M'
                OR berth LIKE '%%WR%%19%%'
            )
            GROUP BY cargo
        )

        SELECT
            berth,
            cargo,
            balance
        FROM final_report
        WHERE balance > 0
        ORDER BY
            CASE
                WHEN berth LIKE '%%10%%' THEN 10
                WHEN berth LIKE '%%11%%' THEN 11
                WHEN berth LIKE '%%12%%' THEN 12
                WHEN berth LIKE '%%19%%' OR berth = 'WR 19' THEN 19
                WHEN berth = 'JETTY' THEN 100
                ELSE 99
            END,
            cargo;
    """

    cur.execute(sql_query, (target_date_str, shift_key, cutoff_date_str))
    rows = cur.fetchall()
    cur.close()
    conn.close()

    table_items = []
    for r in rows:
        r_bal = int(round(float(r['balance'] or 0)))
        if r_bal <= 0:
            continue
        raw_berth = (r.get('berth') or '').strip().upper()
        cargo = _clean_cargo_name(r.get('cargo'))

        if _is_wr19(raw_berth):
            b_display = 'WR 19'
            note = '(WR 19)'
            is_sp = True
        else:
            is_sp, b_num = _is_berth_10_to_12(raw_berth)
            if is_sp:
                b_display = f"BERTH {b_num}"
                note = f"(Berth No.{b_num})"
                is_sp = True
            else:
                b_display = 'Jetty'
                note = ''
                is_sp = False

        table_items.append({
            'berth': b_display,
            'cargo': cargo,
            'balance': r_bal,
            'note': note,
            'is_special': is_sp
        })

    grand_total = sum(it['balance'] for it in table_items)

    # Build WhatsApp / SMS Text Block
    sms_lines = [f"Cargo Balance at Jetty for {shift_key} Shift", ""]
    if table_items:
        max_c_len = max(len(it['cargo']) for it in table_items)
        max_c_len = max(max_c_len, 14)
        for it in table_items:
            c_str = it['cargo'].ljust(max_c_len)
            b_str = f"{it['balance']:,} MT."
            if it['note']:
                sms_lines.append(f"{c_str} : {b_str} {it['note']}")
            else:
                sms_lines.append(f"{c_str} : {b_str}")

        sms_lines.append("")
        sms_lines.append(f"Total: {grand_total:,} MT.")
    else:
        sms_lines.append("No active cargo balance at jetty for this shift.")
        sms_lines.append("")
        sms_lines.append("Total: 0 MT.")

    sms_lines.append("")
    sms_lines.append("Regards")
    sms_text = "\n".join(sms_lines)

    return {
        'entry_date': target_date_str,
        'shift': shift_key,
        'shift_display': f"{shift_key} Shift",
        'title': f"Cargo Balance at Jetty for {shift_key} Shift",
        'items': table_items,
        'total_qty': grand_total,
        'total_balance': grand_total,
        'sms_text': sms_text
    }


def _current_shift_code(now=None):
    now = now or datetime.now()
    hour = now.hour
    if 6 <= hour < 14:
        return 'A'
    if 14 <= hour < 22:
        return 'B'
    return 'C'


def _operational_date(now=None):
    now = now or datetime.now()
    d = now.date()
    if now.hour < 6:
        d -= timedelta(days=1)
    return d


# ── Routes ───────────────────────────────────────────────────────────────────

@bp.route('/module/RP01/shift-cargo-balance/')
@login_required
def shift_cargo_balance_index():
    now = datetime.now()
    cur_op = _operational_date(now)
    if cur_op < REPORT_CUTOFF_DATE:
        cur_op = REPORT_CUTOFF_DATE
    today_str = cur_op.strftime('%Y-%m-%d')
    cur_shift = _current_shift_code(now)
    return render_template(
        'shift_cargo_balance/shift_cargo_balance.html',
        username=session.get('username'),
        default_date=today_str,
        default_shift=cur_shift
    )


@bp.route('/api/module/RP01/shift-cargo-balance/preview')
@bp.route('/api/module/RP01/shift-cargo-balance/data')
@login_required
def shift_cargo_balance_data():
    now = datetime.now()
    entry_date = (request.args.get('entry_date') or request.args.get('date') or '').strip()
    shift = (request.args.get('shift') or '').strip().upper()

    if not entry_date:
        entry_date = _operational_date(now).strftime('%Y-%m-%d')
    if shift not in ('A', 'B', 'C'):
        shift = _current_shift_code(now)

    data = _fetch_shift_cargo_balance(entry_date, shift)
    return jsonify(data)


@bp.route('/api/module/RP01/shift-cargo-balance/download')
@login_required
def shift_cargo_balance_download():
    now = datetime.now()
    entry_date = (request.args.get('entry_date') or request.args.get('date') or '').strip()
    shift = (request.args.get('shift') or '').strip().upper()

    if not entry_date:
        entry_date = _operational_date(now).strftime('%Y-%m-%d')
    if shift not in ('A', 'B', 'C'):
        shift = _current_shift_code(now)

    data = _fetch_shift_cargo_balance(entry_date, shift)

    # ── Build Excel matching exact format: berth, cargo, balance ──
    wb = Workbook()
    ws = wb.active
    ws.title = f"Shift {shift} Balance"
    ws.views.sheetView[0].showGridLines = True

    _thin = Side(style='thin', color='000000')
    border = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)
    font_title = Font(name='Calibri', size=11, bold=True)
    font_header = Font(name='Calibri', size=11, bold=True)
    font_data = Font(name='Calibri', size=11)
    font_bold = Font(name='Calibri', size=11, bold=True)

    # Row 1: Title (Cargo Balance at Jetty for <shift> Shift)
    ws.cell(1, 1, f"Cargo Balance at Jetty for {shift} Shift").font = font_title

    # Row 2: Headers (BERTH, CARGO, BALANCE)
    headers = ['BERTH', 'CARGO', 'BALANCE']
    for ci, h in enumerate(headers, 1):
        cell = ws.cell(2, ci, h)
        cell.font = font_header
        cell.border = border
        cell.alignment = Alignment(
            horizontal='center' if ci == 1 else ('right' if ci == 3 else 'left')
        )

    # Data Rows
    row_idx = 3
    for it in data['items']:
        berth_display = it['berth']
        if it.get('note'):
            berth_display = f"{it['berth']} {it['note']}"
        c_berth = ws.cell(row_idx, 1, berth_display)
        c_cargo = ws.cell(row_idx, 2, it['cargo'])
        c_bal = ws.cell(row_idx, 3, it['balance'])

        c_berth.font = font_data
        c_berth.border = border
        c_berth.alignment = Alignment(horizontal='left')

        c_cargo.font = font_data
        c_cargo.border = border
        c_cargo.alignment = Alignment(horizontal='left')

        c_bal.font = font_data
        c_bal.border = border
        c_bal.number_format = '#,##0'
        c_bal.alignment = Alignment(horizontal='right')

        row_idx += 1

    # Total Row
    c_tot_lbl = ws.cell(row_idx, 1, 'Total')
    c_tot_lbl.font = font_bold
    c_tot_lbl.border = border

    c_empty = ws.cell(row_idx, 2, '')
    c_empty.font = font_bold
    c_empty.border = border

    c_tot_val = ws.cell(row_idx, 3, data['total_qty'])
    c_tot_val.font = font_bold
    c_tot_val.border = border
    c_tot_val.number_format = '#,##0'
    c_tot_val.alignment = Alignment(horizontal='right')

    # Set column widths
    ws.column_dimensions['A'].width = 22
    ws.column_dimensions['B'].width = 34
    ws.column_dimensions['C'].width = 18

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    filename = f"Cargo_Balance_Jetty_{entry_date}_Shift_{shift}.xlsx"
    return Response(
        buf.getvalue(),
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        headers={'Content-Disposition': f'attachment; filename="{filename}"'}
    )
